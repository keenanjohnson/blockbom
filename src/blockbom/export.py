"""Exporters for project-based BOMs: CSV, XLSX, and Mermaid."""

import csv
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .project import Component, Project
from .rollup import (
    ConsolidatedRow,
    IndentedRow,
    Rollup,
    consolidated_bom,
    indented_bom,
    project_rollup,
)
from .stamp import export_stamp

INDENTED_HEADERS = [
    "Level",
    "Qty",
    "Extended Qty",
    "Name",
    "Part Number",
    "Supplier",
    "Supplier PN",
    "Link",
    "Unit Cost",
    "Extended Cost",
    "Unit Weight (g)",
    "Extended Weight (g)",
    "Subtree Cost",
    "Subtree Weight (g)",
]

CONSOLIDATED_HEADERS = [
    "Part Number",
    "Name",
    "Total Qty",
    "Supplier",
    "Supplier PN",
    "Link",
    "Unit Cost",
    "Extended Cost",
    "Unit Weight (g)",
    "Extended Weight (g)",
]


def _fmt(value: float | None) -> str:
    return f"{value:.2f}" if value is not None else ""


def _fmt_rollup(total: float, complete: bool) -> str:
    return f"{total:.2f}" if complete else f">= {total:.2f} (incomplete)"


def _extra_columns(extras: list[dict[str, Any]]) -> list[str]:
    columns: list[str] = []
    for extra in extras:
        for key in extra:
            if key not in columns:
                columns.append(key)
    return columns


def _indented_table(project: Project) -> tuple[list[str], list[list[str]]]:
    rows = indented_bom(project)
    extra_cols = _extra_columns([r.extra for r in rows])
    headers = INDENTED_HEADERS + extra_cols
    table = [
        [*_indented_cells(row), *(str(row.extra.get(c, "")) for c in extra_cols)] for row in rows
    ]
    return headers, table


def _indented_cells(row: IndentedRow) -> list[str]:
    return [
        str(row.level),
        str(row.qty),
        str(row.extended_qty),
        row.name,
        row.part_number,
        row.supplier,
        row.supplier_pn,
        row.link,
        _fmt(row.unit_cost),
        _fmt(row.extended_cost),
        _fmt(row.unit_weight_g),
        _fmt(row.extended_weight_g),
        _fmt_rollup(row.rollup.total_cost, row.rollup.cost_complete) if row.rollup else "",
        _fmt_rollup(row.rollup.total_weight_g, row.rollup.weight_complete) if row.rollup else "",
    ]


def _consolidated_table(project: Project) -> tuple[list[str], list[list[str]]]:
    rows = consolidated_bom(project)
    extra_cols = _extra_columns([r.extra for r in rows])
    headers = CONSOLIDATED_HEADERS + extra_cols
    table = [
        [*_consolidated_cells(row), *(str(row.extra.get(c, "")) for c in extra_cols)]
        for row in rows
    ]
    rollup = project_rollup(project)
    table.append(_totals_row(len(headers), rollup))
    return headers, table


def _consolidated_cells(row: ConsolidatedRow) -> list[str]:
    return [
        row.part_number,
        row.name,
        str(row.total_qty),
        row.supplier,
        row.supplier_pn,
        row.link,
        _fmt(row.unit_cost),
        _fmt(row.extended_cost),
        _fmt(row.unit_weight_g),
        _fmt(row.extended_weight_g),
    ]


def _totals_row(width: int, rollup: Rollup) -> list[str]:
    row = [""] * width
    row[1] = "TOTAL"
    row[7] = _fmt_rollup(rollup.total_cost, rollup.cost_complete)
    row[9] = _fmt_rollup(rollup.total_weight_g, rollup.weight_complete)
    return row


def _write_csv(path: Path | str, headers: list[str], table: list[list[str]], stamp: str) -> None:
    with open(path, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(headers)
        writer.writerows(table)
        writer.writerow([])
        writer.writerow([f"# {stamp}"])


def export_indented_csv(
    project: Project, path: Path | str, source: Path | str | None = None
) -> None:
    """Write the hierarchical BOM as CSV.

    `source` is the project.yaml path; its directory determines the git
    revision in the stamp.
    """
    headers, table = _indented_table(project)
    _write_csv(path, headers, table, export_stamp(project.name, _stamp_cwd(source)))


def export_consolidated_csv(
    project: Project, path: Path | str, source: Path | str | None = None
) -> None:
    """Write the consolidated (flat, per-part) BOM as CSV, with a totals row."""
    headers, table = _consolidated_table(project)
    _write_csv(path, headers, table, export_stamp(project.name, _stamp_cwd(source)))


def _stamp_cwd(source: Path | str | None) -> Path | None:
    return Path(source).resolve().parent if source is not None else None


def export_xlsx(project: Project, path: Path | str, source: Path | str | None = None) -> None:
    """Write an XLSX workbook: Summary, Indented BOM, and Consolidated BOM sheets."""
    workbook = Workbook()
    bold = Font(bold=True)

    summary = workbook.active
    assert summary is not None
    summary.title = "Summary"
    rollup = project_rollup(project)
    summary_rows: list[tuple[str, str]] = [
        ("Project", project.name),
        ("Stamp", export_stamp(project.name, _stamp_cwd(source))),
        ("Total Cost", _fmt_rollup(rollup.total_cost, rollup.cost_complete)),
        ("Total Weight (g)", _fmt_rollup(rollup.total_weight_g, rollup.weight_complete)),
        ("Part Count", str(rollup.part_count)),
    ]
    for label, value in summary_rows:
        summary.append([label, value])
    for cell_row in summary.iter_rows(max_col=1):
        for cell in cell_row:
            cell.font = bold
    summary.column_dimensions["A"].width = 18
    summary.column_dimensions["B"].width = 60

    for title, (headers, table) in (
        ("Indented BOM", _indented_table(project)),
        ("Consolidated BOM", _consolidated_table(project)),
    ):
        sheet = workbook.create_sheet(title)
        _fill_sheet(sheet, headers, table, bold)

    workbook.save(str(path))


# Columns whose values may safely become numbers in XLSX. Everything else
# (Part Number, Supplier PN, ...) stays text: "0402" must never become 402.
NUMERIC_COLUMNS = frozenset(
    {
        "Level",
        "Qty",
        "Extended Qty",
        "Total Qty",
        "Unit Cost",
        "Extended Cost",
        "Unit Weight (g)",
        "Extended Weight (g)",
        "Subtree Cost",
        "Subtree Weight (g)",
    }
)


def _fill_sheet(sheet: Worksheet, headers: list[str], table: list[list[str]], bold: Font) -> None:
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = bold
    for row in table:
        sheet.append(_typed_cells(row, headers))
    for index, header in enumerate(headers):
        width = max([len(header)] + [len(str(row[index])) for row in table if index < len(row)])
        column = get_column_letter(index + 1)
        sheet.column_dimensions[column].width = min(width + 2, 50)
    sheet.freeze_panes = "A2"


def _typed_cells(row: list[str], headers: list[str]) -> list[Any]:
    """Convert numeric-column cells to numbers so Excel can sum them."""
    typed: list[Any] = []
    for index, value in enumerate(row):
        header = headers[index] if index < len(headers) else ""
        if header in NUMERIC_COLUMNS:
            try:
                number = float(value)
                typed.append(int(number) if number == int(number) else number)
                continue
            except ValueError:
                pass
        typed.append(value)
    return typed


def emit_mermaid(project: Project) -> str:
    """Render the project as a Mermaid flowchart."""
    lines = ["flowchart TD"]

    def label(component: Component) -> str:
        text = component.name.replace('"', "'")
        if component.qty > 1:
            text = f"{component.qty}x {text}"
        return text

    def visit(component: Component, indent: int) -> None:
        pad = "    " * indent
        if component.children:
            lines.append(f'{pad}subgraph {component.id}["{label(component)}"]')
            for child in component.children:
                visit(child, indent + 1)
            lines.append(f"{pad}end")
        else:
            lines.append(f'{pad}{component.id}["{label(component)}"]')

    for component in project.components:
        visit(component, 1)

    for connection in project.connections:
        if connection.label:
            lines.append(f"    {connection.source} -->|{connection.label}| {connection.target}")
        else:
            lines.append(f"    {connection.source} --> {connection.target}")

    return "\n".join(lines) + "\n"


def export_mermaid(project: Project, path: Path | str) -> None:
    """Write the project's block diagram as a Mermaid .mmd file."""
    Path(path).write_text(emit_mermaid(project), encoding="utf-8")
