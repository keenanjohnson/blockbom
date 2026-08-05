"""Tests for the project-based exporters (CSV, XLSX, Mermaid, stamp)."""

import csv

from openpyxl import load_workbook

from blockbom import parse_mermaid
from blockbom.export import (
    emit_mermaid,
    export_consolidated_csv,
    export_indented_csv,
    export_xlsx,
)
from blockbom.project import Component, Connection, Part, Project
from blockbom.stamp import export_stamp, git_revision


def _project() -> Project:
    return Project(
        name="Widget",
        components=[
            Component(
                id="board",
                name="Driver Board",
                qty=2,
                children=[
                    Component(id="mcu", name="MCU", part="MCU-1"),
                    Component(id="screws", name="Screws", part="0402", qty=4),
                ],
            ),
            Component(id="display", name="Display", part="DISP-1"),
        ],
        connections=[Connection.model_validate({"from": "mcu", "to": "display", "label": "SPI"})],
        parts={
            "MCU-1": Part(cost=5.00, weight_g=10, supplier="Digi-Key", mfr="ST"),
            "0402": Part(description="Passive", cost=0.10, weight_g=1),
            "DISP-1": Part(cost=20.00, supplier="Adafruit"),
        },
    )


def _read_csv(path):
    with open(path, newline="", encoding="utf-8") as f:
        return list(csv.reader(f))


class TestIndentedCSV:
    def test_rows_and_rollups(self, tmp_path):
        path = tmp_path / "bom.csv"
        export_indented_csv(_project(), path)
        rows = _read_csv(path)
        headers = rows[0]
        assert "Subtree Cost" in headers
        board = next(r for r in rows if r[headers.index("Name")] == "Driver Board")
        # 2 x (5.00 + 4 * 0.10) = 10.80
        assert board[headers.index("Subtree Cost")] == "10.80"
        screws = next(r for r in rows if r[headers.index("Name")] == "Screws")
        assert screws[headers.index("Extended Qty")] == "8"
        assert screws[headers.index("Extended Cost")] == "0.80"

    def test_extra_part_fields_become_columns(self, tmp_path):
        path = tmp_path / "bom.csv"
        export_indented_csv(_project(), path)
        rows = _read_csv(path)
        headers = rows[0]
        assert "mfr" in headers
        mcu = next(r for r in rows if r[headers.index("Name")] == "MCU")
        assert mcu[headers.index("mfr")] == "ST"

    def test_stamp_footer(self, tmp_path):
        path = tmp_path / "bom.csv"
        export_indented_csv(_project(), path)
        rows = _read_csv(path)
        assert rows[-1][0].startswith("# Widget | exported ")

    def test_incomplete_weight_marked(self, tmp_path):
        # DISP-1 has no weight -> project weight rollup incomplete
        path = tmp_path / "flat.csv"
        export_consolidated_csv(_project(), path)
        rows = _read_csv(path)
        headers = rows[0]
        total = next(r for r in rows if len(r) > 1 and r[1] == "TOTAL")
        assert total[headers.index("Extended Cost")] == "30.80"
        weight = total[headers.index("Extended Weight (g)")]
        assert weight.startswith(">=") and "incomplete" in weight


class TestConsolidatedCSV:
    def test_grouped_by_part(self, tmp_path):
        path = tmp_path / "flat.csv"
        export_consolidated_csv(_project(), path)
        rows = _read_csv(path)
        headers = rows[0]
        screws = next(r for r in rows if r[headers.index("Part Number")] == "0402")
        assert screws[headers.index("Total Qty")] == "8"


class TestXLSX:
    def test_workbook_sheets_and_values(self, tmp_path):
        path = tmp_path / "bom.xlsx"
        export_xlsx(_project(), path)
        workbook = load_workbook(path)
        assert workbook.sheetnames == ["Summary", "Indented BOM", "Consolidated BOM"]
        summary = {row[0].value: row[1].value for row in workbook["Summary"].iter_rows()}
        assert summary["Project"] == "Widget"
        assert summary["Total Cost"] == "30.80"

    def test_part_numbers_stay_text_but_costs_are_numeric(self, tmp_path):
        path = tmp_path / "bom.xlsx"
        export_xlsx(_project(), path)
        sheet = load_workbook(path)["Consolidated BOM"]
        headers = [c.value for c in sheet[1]]
        pn_col = headers.index("Part Number")
        cost_col = headers.index("Unit Cost")
        values = {row[pn_col].value for row in sheet.iter_rows(min_row=2)}
        assert "0402" in values  # not 402
        screws_row = next(
            row for row in sheet.iter_rows(min_row=2) if row[pn_col].value == "0402"
        )
        assert screws_row[cost_col].value == 0.10


class TestMermaidEmit:
    def test_emit_structure(self):
        text = emit_mermaid(_project())
        assert text.startswith("flowchart TD")
        assert 'subgraph board["2x Driver Board"]' in text
        assert 'screws["4x Screws"]' in text
        assert "mcu -->|SPI| display" in text

    def test_emitted_mermaid_reparses(self):
        text = emit_mermaid(_project())
        diagram = parse_mermaid(text)
        assert "mcu" in diagram.nodes
        assert diagram.subgraphs[0].id == "board"
        assert any(e.source_id == "mcu" and e.target_id == "display" for e in diagram.edges)


class TestStamp:
    def test_git_revision_in_repo(self):
        revision = git_revision()
        assert revision is not None

    def test_stamp_format(self):
        stamp = export_stamp("Widget")
        assert stamp.startswith("Widget | exported 20")
        assert "git" in stamp

    def test_no_git_outside_repo(self, tmp_path):
        assert git_revision(tmp_path) is None
